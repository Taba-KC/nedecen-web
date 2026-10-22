from pathlib import Path
from zipfile import ZipFile
import json,hashlib,re
from lxml import etree
from docx import Document
root=Path(__file__).resolve().parent
repo=root.parent.parent
manifest=json.loads((root/'.qa/source_manifest.json').read_text(encoding='utf-8'))
changed=[]
for f in manifest['files']:
 p=repo.parent/f['repository']/f['path']
 if hashlib.sha256(p.read_bytes()).hexdigest()!=f['sha256']:changed.append(f['path'])
doc=Document(root/'NEDECEN Development Journey.docx')
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
assert '\u2014' not in text
assert not changed,changed
assert 'ac331ba' in text and 'c3ad101' in text
assert 'Chapter 1 Building the backend' in text
assert 'Chapter 2 Building the frontend' in text
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(root/'NEDECEN Development Journey.docx') as z:
 tree=etree.fromstring(z.read('word/styles.xml'))
 borders=tree.xpath('//w:style[@w:styleId="Title"]//w:pBdr',namespaces=ns)
 assert not borders
 tree=etree.fromstring(z.read('word/document.xml'))
 titles=tree.xpath('//w:p[w:pPr/w:pStyle[@w:val="Title"]]',namespaces=ns)
 assert len(titles)==1
 for e in titles[0].xpath('.//w:pBdr/*',namespaces=ns):assert e.get('{'+ns['w']+'}val') in ['nil','none']
 toc_fields=tree.xpath('//w:instrText[contains(text(),"TOC")]',namespaces=ns)
 assert toc_fields,'Word did not materialize the contents field'
report={'source_files_verified':len(manifest['files']),'original_source_lines':manifest['source_lines'],'database_tables':manifest['database_tables'],'word_tables':len(doc.tables),'em_dashes':0,'source_files_changed_since_snapshot':changed,'contents_updated':True,'title_border_removed':True}
(root/'.qa/validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
